"""Genera el documento de evidencias del taller (Word y Markdown).

    uv run python documento/generar_documento.py [--autor "Nombre"]

Para cada punto reúne:
  1. La evidencia de ejecución que dejaron los scripts (evidencias/puntoN/log.md y sus tablas).
  2. Las capturas de pantalla que usted guarde en evidencias/puntoN/ (.png / .jpg).
  3. Las respuestas teóricas y de costos (documento/RESPUESTAS_TEORICAS.md).
  4. Un anexo con la lista de archivos de evidencia (JSON, CSV, salidas de Terraform/CLI).

Salida: documento/EVIDENCIAS_Actividad_AWS.docx y documento/EVIDENCIAS_Actividad_AWS.md
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
EVID = RAIZ / "evidencias"
SALIDA_DOCX = RAIZ / "documento" / "EVIDENCIAS_Actividad_AWS.docx"
SALIDA_MD = RAIZ / "documento" / "EVIDENCIAS_Actividad_AWS.md"

PUNTOS = {
    1: ("AWS Secrets Manager", "punto1_secrets_manager"),
    2: ("Despliegue con Terraform, SDK y AWS CLI", "punto2_terraform_sdk_cli"),
    3: ("Amazon DocumentDB", "punto3_documentdb"),
    4: ("Amazon Neptune", "punto4_neptune"),
    5: ("Amazon Redshift", "punto5_redshift"),
    6: ("AWS Glue Data Catalog", "punto6_glue_catalog"),
}
IMAGENES = {".png", ".jpg", ".jpeg"}
MAX_FILAS_TABLA = 40
AZUL = RGBColor(0x1F, 0x3A, 0x5F)


# ----------------------------------------------------------------- Markdown -> Word
def _runs(parrafo, texto: str) -> None:
    """Agrega texto con **negrita**, `código` y _cursiva_ simples."""
    partes = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\w)_[^_]+_(?!\w))", texto)
    for p in partes:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            parrafo.add_run(p[2:-2]).bold = True
        elif p.startswith("`") and p.endswith("`"):
            r = parrafo.add_run(p[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif p.startswith("_") and p.endswith("_") and len(p) > 2:
            parrafo.add_run(p[1:-1]).italic = True
        else:
            parrafo.add_run(p)


def _sombrear(celda, color: str) -> None:
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc.append(shd)


def _tabla(doc, filas: list[list[str]]) -> None:
    if not filas:
        return
    total = len(filas) - 1
    if total > MAX_FILAS_TABLA:
        filas = filas[:MAX_FILAS_TABLA + 1]
    ncols = max(len(f) for f in filas)
    t = doc.add_table(rows=len(filas), cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j in range(ncols):
            celda = t.cell(i, j)
            celda.text = ""
            p = celda.paragraphs[0]
            _runs(p, fila[j] if j < len(fila) else "")
            for r in p.runs:
                r.font.size = Pt(8.5)
                if i == 0:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                _sombrear(celda, "1F3A5F")
    if total > MAX_FILAS_TABLA:
        doc.add_paragraph(f"({total} filas en total; se muestran {MAX_FILAS_TABLA}. "
                          "Ver el CSV en la carpeta de evidencias.)").runs[0].italic = True
    doc.add_paragraph()


def _celdas(linea: str) -> list[str]:
    return [c.strip() for c in linea.strip().strip("|").split("|")]


def md_a_docx(doc, texto: str, nivel_base: int = 2) -> None:
    lineas = texto.splitlines()
    i = 0
    while i < len(lineas):
        ln = lineas[i]
        s = ln.strip()
        if s.startswith("```"):
            lenguaje = s[3:].strip()
            bloque = []
            i += 1
            while i < len(lineas) and not lineas[i].strip().startswith("```"):
                bloque.append(lineas[i])
                i += 1
            if lenguaje == "mermaid":
                doc.add_paragraph("Diagrama (Mermaid; se visualiza en el README de GitHub):").runs[0].italic = True
            p = doc.add_paragraph()
            r = p.add_run("\n".join(bloque))
            r.font.name = "Consolas"
            r.font.size = Pt(8)
            p.paragraph_format.left_indent = Cm(0.5)
        elif s.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", lineas[i].strip()):
                    filas.append(_celdas(lineas[i]))
                i += 1
            _tabla(doc, filas)
            continue
        elif s.startswith("#"):
            n = len(s) - len(s.lstrip("#"))
            doc.add_heading(s[n:].strip(), level=min(nivel_base + n - 1, 4))
        elif s in ("---", "***"):
            pass
        elif re.match(r"^([-*]|\d+\.) ", s):
            numerada = s[0].isdigit()
            texto = s if numerada else s[2:]
            # une las líneas de continuación (texto con sangría)
            while i + 1 < len(lineas) and lineas[i + 1].startswith("  ") and lineas[i + 1].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.) ", lineas[i + 1]):
                i += 1
                texto += " " + lineas[i].strip()
            if numerada:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.4)
            else:
                p = doc.add_paragraph(style="List Bullet")
            _runs(p, texto)
        elif s:
            # párrafo: une líneas consecutivas de texto normal
            texto = s
            while i + 1 < len(lineas):
                sig = lineas[i + 1].strip()
                if not sig or re.match(r"^(#|\||```|[-*] |\d+\. |---)", sig):
                    break
                texto += " " + sig
                i += 1
            _runs(doc.add_paragraph(), texto)
        i += 1


# ----------------------------------------------------------------- contenido
def respuestas_por_punto() -> dict[int, str]:
    archivo = RAIZ / "documento" / "RESPUESTAS_TEORICAS.md"
    if not archivo.exists():
        return {}
    bloques, actual, buf = {}, None, []
    for ln in archivo.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^## Punto (\d)", ln)
        if m:
            if actual:
                bloques[actual] = "\n".join(buf)
            actual, buf = int(m.group(1)), []
        elif actual:
            buf.append(ln)
    if actual:
        bloques[actual] = "\n".join(buf)
    return bloques


def log_limpio(punto: int) -> str:
    log = EVID / f"punto{punto}" / "log.md"
    if not log.exists():
        return ""
    lineas = []
    for ln in log.read_text(encoding="utf-8").splitlines():
        if ln.startswith("# Evidencias"):
            continue
        if re.match(r"^- (Evidencia|Archivo): `", ln):
            continue  # se listan en el anexo
        lineas.append(ln)
    return "\n".join(lineas)


def archivos_evidencia(punto: int) -> list[Path]:
    d = EVID / f"punto{punto}"
    if not d.exists():
        return []
    return sorted(p for p in d.rglob("*") if p.is_file() and p.name != ".gitkeep")


def generar(autor: str) -> None:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)

    # Portada
    for _ in range(6):
        doc.add_paragraph()
    for texto, tam, negrita in (("Universidad EIA", 14, False), ("Ingeniería de Datos 2026-2", 14, False),
                                ("Actividad AWS", 28, True), ("Documento de evidencias", 16, False),
                                ("", 12, False), (autor, 13, False),
                                (dt.date.today().strftime("%d/%m/%Y"), 12, False)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(texto)
        r.font.size = Pt(tam)
        r.bold = negrita
        if negrita:
            r.font.color.rgb = AZUL
    doc.add_page_break()

    doc.add_heading("Contenido", level=1)
    for n, (titulo, carpeta) in PUNTOS.items():
        doc.add_paragraph(f"Punto {n} – {titulo}  (código: {carpeta}/)", style="List Number")
    doc.add_paragraph("Región: us-east-1. Todos los recursos se crearon con el usuario IAM de trabajo "
                      "(no la cuenta raíz) y se eliminaron al terminar cada punto; la verificación de "
                      "eliminación aparece al final de la evidencia de cada punto. Ningún valor de secreto "
                      "ni clave de acceso aparece en este documento.")
    doc.add_page_break()

    respuestas = respuestas_por_punto()
    md = [f"# Actividad AWS – Documento de evidencias\n\n{autor} – {dt.date.today():%d/%m/%Y}\n"]
    for n, (titulo, carpeta) in PUNTOS.items():
        doc.add_heading(f"Punto {n} – {titulo}", level=1)
        md.append(f"\n# Punto {n} – {titulo}\n")
        p = doc.add_paragraph()
        _runs(p, f"Código y README: `{carpeta}/`. Evidencia automática: `evidencias/punto{n}/`.")

        doc.add_heading("Evidencia de ejecución", level=2)
        md.append("\n## Evidencia de ejecución\n")
        log = log_limpio(n)
        if log.strip():
            md_a_docx(doc, log, nivel_base=2)
            md.append(log)
        else:
            doc.add_paragraph("(Pendiente: ejecute el script de este punto para generar la evidencia.)")
            md.append("_(Pendiente de ejecución)_")

        imagenes = [f for f in archivos_evidencia(n) if f.suffix.lower() in IMAGENES]
        if imagenes:
            doc.add_heading("Capturas de pantalla", level=2)
            md.append("\n## Capturas de pantalla\n")
            for img in imagenes:
                try:
                    doc.add_picture(str(img), width=Cm(16))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap = doc.add_paragraph(img.stem.replace("_", " "))
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.runs[0].italic = True
                    md.append(f"![{img.stem}](../{img.relative_to(RAIZ).as_posix()})")
                except Exception as e:  # noqa: BLE001
                    doc.add_paragraph(f"(No se pudo insertar {img.name}: {e})")

        if n in respuestas:
            doc.add_heading("Respuestas teóricas y de costos", level=2)
            md_a_docx(doc, respuestas[n], nivel_base=2)
            md.append("\n## Respuestas teóricas y de costos\n" + respuestas[n])

        archivos = [f for f in archivos_evidencia(n) if f.suffix.lower() not in IMAGENES]
        if archivos:
            doc.add_heading("Archivos de evidencia", level=2)
            for f in archivos:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(f.relative_to(RAIZ).as_posix())
                r.font.size = Pt(8.5)
                r.font.name = "Consolas"
        doc.add_page_break()

    doc.save(SALIDA_DOCX)
    SALIDA_MD.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"Generado: {SALIDA_DOCX.relative_to(RAIZ)}")
    print(f"Generado: {SALIDA_MD.relative_to(RAIZ)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--autor", default="Luis Miguel García Castaño y Julio José Jaller Cordero")
    generar(ap.parse_args().autor)
